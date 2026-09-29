from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = '攻击动作测试用例'

header_font = Font(bold=True, size=11, color='FFFFFF')
header_fill = PatternFill('solid', start_color='2F75B5')
center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)
p0_fill = PatternFill('solid', start_color='FF6B6B')
p1_fill = PatternFill('solid', start_color='FFA94D')
p2_fill = PatternFill('solid', start_color='FFE066')
p3_fill = PatternFill('solid', start_color='69DB7C')

headers = ['用例ID', '模块', '测试项', '前置条件', '操作步骤', '预期结果', '优先级', '实际结果', '是否通过', '备注']
ws.append(headers)
for col, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center_align
    cell.border = thin_border

cases = [
    ['TC-ATT-001', '基础攻击', '待机→普攻过渡', '角色处于待机状态，无异常状态',
     '1. 角色静止站立\n2. 按下普攻键一次\n3. 观察从待机到攻击的切换过程',
     '待机到普攻之间应有平滑过渡动画（Blend），不应出现瞬间跳变或画面闪烁', 'P0', '', '', '核心问题：僵硬感主要来源'],
    ['TC-ATT-002', '基础攻击', '普攻→待机恢复', '角色正在执行普攻动作',
     '1. 执行一次完整普攻\n2. 攻击结束后等待角色恢复\n3. 观察攻击结束到待机的过渡',
     '攻击结束后应平滑过渡回待机姿态，收招动作完整播放，无突然卡回待机的情况', 'P0', '', '', '核心问题：僵硬感主要来源'],
    ['TC-ATT-003', '基础攻击', '普攻动作完整性', '角色可正常执行攻击',
     '1. 按下普攻键\n2. 完整观察整个攻击动画（起手→挥击→收招）',
     '攻击动画应完整播放：包含起势、攻击、收招三个阶段，无中途被截断或跳帧现象', 'P0', '', '', '检查Animation Clip时长设置'],
    ['TC-ATT-004', '基础攻击', '普攻节奏/速度感', '角色可正常执行攻击',
     '1. 连续执行5次普攻\n2. 感受每次攻击的节奏和速度',
     '攻击动作应有合理的加速和减速曲线（缓入缓出），不应全程匀速或机械式重复', 'P1', '', '', '检查Animation Curve'],
    ['TC-ATT-005', '连招系统', '二段连击衔接', '角色可正常执行普攻，已配置连招',
     '1. 执行第一段普攻\n2. 在合适时间窗口内按下第二次普攻\n3. 观察两段攻击之间的衔接',
     '第一段攻击应自然衔接第二段攻击，中间有过渡Blend或专用衔接动画，无明显停顿或跳变', 'P0', '', '', '重点测试项'],
    ['TC-ATT-006', '连招系统', '三段连击衔接', '同上',
     '1. 执行一段→二段→三段连续攻击\n2. 逐段观察衔接流畅度',
     '三段连击应形成连贯的攻击组合，每段间衔接平滑，整体观感一气呵成', 'P0', '', '', '重点测试项'],
    ['TC-ATT-007', '连招系统', '连招输入窗口', '角色可正常执行攻击',
     '1. 第一段攻击后立即（<0.1s）按第二段\n2. 第一段攻击后在适当时机（0.1~0.5s）按第二段\n3. 第一段攻击后超时（>1s）按第二段',
     '在输入窗口内应成功触发连击；过早或过晚输入不应触发连击或应有合理处理', 'P1', '', '', '测试Animator Transition条件'],
    ['TC-ATT-008', '连招系统', '连招中断处理', '角色正在执行连招',
     '1. 执行连招过程中按防御/闪避键\n2. 连招过程中受到敌人攻击\n3. 观察中断后的表现',
     '连招被中断时应平滑切换到对应状态（防御/受击），不应出现T-pose或动作错乱', 'P1', '', '', '检查Cancel配置'],
    ['TC-ATT-009', '移动攻击', '跑动中普攻', '角色处于跑动状态',
     '1. 角色向前跑动\n2. 跑动中按下普攻键\n3. 观察攻击动作与移动的配合',
     '跑动中攻击应自然衔接，角色不应出现脚底打滑、位置瞬间重置、或身体与移动方向不一致', 'P0', '', '', 'Root Motion相关关键测试'],
    ['TC-ATT-010', '移动攻击', '各方向移动中攻击', '角色可多方向移动',
     '1. 分别在前/后/左/右/斜向移动时攻击\n2. 观察各方向攻击表现',
     '所有方向的移动攻击均应协调一致，无某个方向特别僵硬或不自然的情况', 'P1', '', '', '方向性全面覆盖'],
    ['TC-ATT-011', '移动攻击', '攻击后恢复移动', '角色在移动中执行了攻击',
     '1. 移动中攻击\n2. 攻击结束后继续移动\n3. 观察攻击→移动的过渡',
     '攻击结束后应无缝恢复移动状态，无停顿、无方向突变、无速度异常', 'P1', '', '', ''],
    ['TC-ATT-012', '方向控制', '攻击朝向跟随', '角色周围有敌人目标',
     '1. 面对敌人攻击\n2. 侧向面对敌人攻击\n3. 背对敌人时攻击\n4. 观察每次攻击的朝向',
     '攻击方向应与角色朝向/目标方向一致，不应出现攻击方向与面向不符的情况', 'P0', '', '', 'LookAt/旋转逻辑测试'],
    ['TC-ATT-013', '方向控制', '攻击中的转向', '角色正在执行攻击动作',
     '1. 执行攻击过程中转动鼠标/摇杆\n2. 观察角色是否能转向及转向方式',
     '根据设计需求验证：若允许攻击中转向则应平滑旋转；若不允许则保持当前朝向直到攻击结束', 'P1', '', '', '需确认设计意图'],
    ['TC-ATT-014', '方向控制', '360度全方位攻击', '角色可自由操控朝向',
     '1. 分别在0°/45°/90°/135°/180°/225°/270°/315°八个方向执行攻击\n2. 对比各方向攻击表现',
     '360度各方向攻击动作表现应一致，无特定角度动作变形或异常', 'P2', '', '', '全角度覆盖'],
    ['TC-ATT-015', '技能攻击', '主动技能动作', '角色已学习至少一个主动技能',
     '1. 按下主动技能按键\n2. 完整观察技能释放的全过程',
     '技能应有完整的施法前摇→释放动作→后摇，动作流畅不卡顿，特效与动作同步', 'P0', '', '', ''],
    ['TC-ATT-016', '技能攻击', '技能取消/打断', '角色正在释放有前摇的技能',
     '1. 技能前摇期间按闪避/移动\n2. 技能释放中受到攻击\n3. 观察被打断时的表现',
     '技能被中断时应正确进入对应状态（闪避/受击），无动作残留或状态卡死', 'P1', '', '', ''],
    ['TC-ATT-017', '技能攻击', '技能衔接普攻', '角色刚释放完一个技能',
     '1. 释放技能后立即按普攻\n2. 观察技能→普攻的衔接',
     '技能结束后接普攻应自然流畅，无额外停顿或异常过渡', 'P1', '', '', ''],
    ['TC-ATT-018', '技能攻击', '连续技能释放', '角色有多个可用技能且CD就绪',
     '1. 快速交替释放不同技能\n2. 观察技能之间的切换',
     '多个技能连续释放时各自动画应正确播放，无动画混乱或状态错误', 'P1', '', '', ''],
    ['TC-ATT-019', '特殊状态', '落地即时攻击', '角色从空中落地瞬间',
     '1. 跳跃后落地的同时按攻击\n2. 观察落地攻击的表现',
     '落地攻击应顺畅执行，落地缓冲与攻击起手应自然融合', 'P2', '', '', ''],
    ['TC-ATT-020', '特殊状态', '受击后反击(霸体/硬直)', '角色刚受到敌人攻击',
     '1. 受到攻击硬直中尝试按攻击\n2. 硬直结束后立即攻击\n3. 观察两种情况',
     '硬直中不应能攻击（除非有霸体机制）；硬直结束后攻击应正常响应', 'P1', '', '', ''],
    ['TC-ATT-021', '特殊状态', '低体力/异常状态下攻击', '角色体力值较低或有debuff',
     '1. 在低体力/debuff状态下执行攻击\n2. 观察动作是否有变化',
     '根据设计确认：低体力时攻击是否应变慢；debuff是否影响攻击动作表现', 'P2', '', '', '需确认具体游戏机制'],
    ['TC-ATT-022', '动画质量', '身体联动(secondary motion)', '角色执行攻击动作',
     '1. 执行攻击时仔细观察身体各部位\n2. 关注：头发/衣物摆动、非武器手臂的姿态、躯干扭转',
     '攻击时身体各部位应有合理的联动反应，不应只有手臂在动而身体其他部位完全静止', 'P1', '', '', '僵硬感的核心原因之一'],
    ['TC-ATT-023', '动画质量', '脚步与地面贴合', '角色执行移动攻击',
     '1. 移动中攻击时聚焦角色脚部\n2. 观察脚底与地面的关系',
     '角色脚底应始终贴合地面，无悬空、穿模、打滑现象', 'P1', '', '', 'IK(反向动力学)相关'],
    ['TC-ATT-024', '动画质量', '武器跟随', '角色持武器执行攻击',
     '1. 攻击过程中观察武器位置\n2. 特别注意快速转身或大动作时',
     '武器应始终紧贴手持部位，无穿模、漂浮或延迟跟随', 'P1', '', '', ''],
    ['TC-ATT-025', '动画质量', '面部表情变化', '角色执行攻击（第三人称视角）',
     '1. 攻击时观察角色面部\n2. 对比待机和攻击中的表情差异',
     '攻击时面部应有对应的用力表情变化，不应全程面无表情（如果项目有做面部动画）', 'P3', '', '', '视项目需求而定'],
    ['TC-ATT-026', '性能测试', '高频连点攻击', '角色可正常操作',
     '1. 以最快速度连续点击攻击键10-20次\n2. 观察动画播放和系统响应',
     '高频输入不应导致动画错乱、帧率骤降、或状态机卡死；系统应正确处理溢出输入', 'P1', '', '', '压力测试'],
    ['TC-ATT-027', '性能测试', '多角色同时攻击', '场景中有多个角色同时进行攻击',
     '1. 让场景中3个以上AI角色/玩家同时攻击\n2. 整体观察表现',
     '多角色同时攻击时各自动画应正常播放，无明显掉帧或动画丢失', 'P2', '', '', '性能压力测试'],
    ['TC-ATT-028', '边界情况', '攻击中场景切换/交互', '角色正在执行攻击动作',
     '1. 攻击过程中触发过图/拾取/对话等\n2. 观察结果',
     '攻击中被场景切换或交互中断时，应正确处理状态清理，无残留动画或状态', 'P2', '', '', ''],
]

priority_map = {'P0': p0_fill, 'P1': p1_fill, 'P2': p2_fill, 'P3': p3_fill}
for row_idx, case in enumerate(cases, 2):
    for col_idx, value in enumerate(case, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.border = thin_border
        if col_idx == 7:
            cell.fill = priority_map.get(value, PatternFill())
            cell.alignment = center_align
            cell.font = Font(bold=True)
        elif col_idx in [1, 2, 7, 9, 10]:
            cell.alignment = center_align
        else:
            cell.alignment = left_align

col_widths = [12, 12, 22, 28, 40, 50, 8, 20, 10, 25]
for i, width in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = width

ws.row_dimensions[1].height = 25
for row in range(2, len(cases) + 2):
    ws.row_dimensions[row].height = 65
ws.freeze_panes = 'A2'

# Sheet2: 测试说明
ws2 = wb.create_sheet('测试说明')
instructions = [
    ['《山海之巅》攻击动作测试用例 — 使用说明'], [''],
    ['一、文档信息'],
    ['项目名称', '山海之巅'],
    ['引擎', 'Unity (3D动作游戏)'],
    ['测试模块', '角色攻击动作系统'],
    ['用例数量', '28条'],
    ['版本', 'v1.0'], [''],
    ['二、优先级说明'],
    ['P0 - 严重（红色）', '核心体验问题，必须修复后才可提测。直接影响攻击手感。'],
    ['P1 - 高（橙色）', '重要问题，强烈建议本迭代修复。影响动作质量和玩家体验。'],
    ['P2 - 中（黄色）', '一般问题，建议修复。在有时间的情况下优先处理。'],
    ['P3 - 低（绿色）', '轻微问题/建议优化。可排入后续迭代。'], [''],
    ['三、重点关注（本次主要问题：攻击僵硬不自然）'],
    ['推荐优先测试以下用例：'],
    ['  TC-ATT-001 待机->普攻过渡', '最可能的僵硬来源：缺少Blend过渡'],
    ['  TC-ATT-002 普攻->待机恢复', '收招动作被截断或无过渡'],
    ['  TC-ATT-005/006 连击衔接', '连招间缺少过渡动画'],
    ['  TC-ATT-009 跑动中攻击', 'Root Motion配置问题导致滑步'],
    ['  TC-ATT-022 身体联动', 'Secondary Motion缺失导致僵硬感'], [''],
    ['四、常见问题排查指引'],
    ['问题现象', '可能原因', '排查方向'],
    ['动作瞬间跳变', 'Animator Transition未配/过渡时间太短', '检查Animator Controller中State间的Transition'],
    ['攻击动作机械重复', '缺少Animation Blend Tree或Curve调优', '检查Animation窗口的Curve编辑器'],
    ['脚底打滑', 'Root Motion未开启或动画未烘焙位移', '检查Animator的Root Motion选项'],
    ['只有手臂在动', '缺少Secondary Animation或骨骼权重问题', '检查动画资源本身的骨骼动画'],
    ['攻击被截断', 'Animation Clip时长不足或State提前退出', '检查Animator State的退出条件和Clip长度'], [''],
    ['五、填写规范'],
    ['实际结果', '如实记录观察到的现象，描述要具体客观'],
    ['是否通过', '填写"通过"/"不通过"/"阻塞"'],
    ['备注', '记录复现步骤、截图路径、视频文件名等附加信息'], [''],
    ['六、Unity Animator相关检查清单（提供给开发参考）'],
    ['Attack State的Transition是否有Exit Time和Transition Duration (>0.1s推荐)'],
    ['是否使用了Blend Tree来混合不同变体的攻击动画'],
    ['Animator是否启用了Root Motion'],
    ['Animation Clip的Curve是否有合理的缓入缓出'],
    ['是否配置了Attack->AnyState的Cancel transition'],
    ['连招各段的Transiton Conditions是否正确配置HasExitTime+参数判断'],
]
for row_data in instructions:
    ws2.append(row_data)

ws2.column_dimensions['A'].width = 40
ws2.column_dimensions['B'].width = 45
ws2.column_dimensions['C'].width = 50
ws2.merge_cells('A1:E1')
ws2['A1'].font = Font(bold=True, size=16, color='2F75B5')
ws2['A1'].alignment = Alignment(horizontal='center')

title_rows = [3, 9, 15, 23, 30, 36]
for r in title_rows:
    if r <= len(instructions):
        ws2.cell(row=r, column=1).font = Font(bold=True, size=12, color='2F75B5')

for col in range(1, 4):
    for r in [24, 37]:
        cell = ws2.cell(row=r, column=col)
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', start_color='D6DCE4')
        cell.border = thin_border

wb.save(r'c:\Users\Administrator\WorkBuddy\Claw\山海之巅_攻击动作_测试用例.xlsx')
print('Done! 28 test cases generated.')
